#!/usr/bin/env python3
"""Regression test for the faithful Excel converter (xlsx_converter.py).

Builds a synthetic workbook with the same structures as a real planning
report (totals above the header, a week-label row over a date header row,
an active autofilter, hidden rows / columns / sheets, accounting number
formats, comments, data validation, merged cells and a pivot table), then
converts it through the given command and checks the Markdown.

    python tests/xlsx_test.py                              # the skill's convert.py
    python tests/xlsx_test.py --cmd dist/markitdown.exe    # the Windows build
"""

from __future__ import annotations

import argparse
import shlex
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_CMD = [sys.executable, str(ROOT / "skills" / "markitdown" / "scripts" / "convert.py")]

ACCOUNTING = '_-* #,##0_-;\\-* #,##0_-;_-* "-"??_-;_-@_-'
RED_PARENS = "#,##0_);[Red](#,##0)"

PIVOT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
  name="PivotFixture" cacheId="1" dataCaption="值">
  <location ref="A3:B6" firstHeaderRow="1" firstDataRow="1" firstDataCol="1" rowPageCount="1" colPageCount="1"/>
  <rowFields count="1"><field x="0"/></rowFields>
  <pageFields count="1"><pageField fld="1" hier="-1"/></pageFields>
  <dataFields count="1"><dataField name="加總 - Qty" fld="2" baseField="0" baseItem="0"/></dataFields>
</pivotTableDefinition>"""
PIVOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotCacheDefinition" Target="../pivotCache/pivotCacheDefinition1.xml"/>
</Relationships>"""
CACHE_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" refreshedDate="46289.5" recordCount="2">
  <cacheSource type="worksheet"><worksheetSource ref="A14:E20" sheet="Report"/></cacheSource>
  <cacheFields count="3"><cacheField name="Brand"/><cacheField name="Region"/><cacheField name="Qty"/></cacheFields>
</pivotCacheDefinition>"""


X14_VALIDATION = (
    '<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
    'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
    '<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
    '<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>Report!$F$15:$F$17</xm:f></x14:formula1>'
    "<xm:sqref>A3:A10 C3:C5 D3:D5 E3:E5 F3:F5</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>"
)


def build_workbook(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    # Rows 1-12: totals and parameters above the real header.
    ws["A1"] = "TTL"
    ws["C1"] = 1234567.891
    ws["C1"].number_format = "#,##0"
    for r in range(2, 13):
        ws.cell(r, 1, f"note {r}")
    # Row 13: week labels over the dates in row 14.
    ws["A13"] = "Average Week"
    ws["D13"] = "W1(W38)"
    ws["E13"] = "W2(W39)"
    for col, name in enumerate(["MPN", "Brand", "Qty"], start=1):
        ws.cell(14, col, name)
    ws["D14"] = 46286  # 2026-09-21
    ws["E14"] = 46293
    for c in ("D14", "E14"):
        ws[c].number_format = "mm/dd/yyyy"
    rows = [
        ("PN-1001", "ACME", 12345678, -1234, 0),
        ("PN-2002", "BETA", 16.399999999999999, 2.069194663831718, 0.5),
        ("HIDDEN-ROW", "BETA", 1, 2, 3),
    ]
    for i, values in enumerate(rows, start=15):
        for col, value in enumerate(values, start=1):
            ws.cell(i, col, value)
        ws.cell(i, 3).number_format = ACCOUNTING
        ws.cell(i, 4).number_format = RED_PARENS
        ws.cell(i, 5).number_format = ACCOUNTING
    ws["C16"].number_format = "General"
    ws["D16"].number_format = "0.00%"
    ws["E15"].number_format = ACCOUNTING  # zero -> "-"
    ws.row_dimensions[17].hidden = True
    ws.column_dimensions["B"].hidden = True
    ws.auto_filter.ref = "A14:E17"
    ws.auto_filter.add_filter_column(0, ["PN-1001", "PN-2002"])
    ws.freeze_panes = "B15"
    ws["C15"].comment = Comment("請確認交期內的訂單是否足夠", "Planner")
    dv = DataValidation(type="list", formula1='"OK,NG,TBD"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("F15:F17")
    ws["F14"] = "Check"
    ws["F15"] = "OK"

    merged = wb.create_sheet("Merged")
    merged["A1"] = "MPN"
    merged["B1"] = "Date"
    merged.merge_cells("B1:B2")
    merged["A2"] = None
    merged["A3"] = "X1"
    merged["B3"] = "2026-W39"
    merged.auto_filter.ref = "A2:B3"

    pivot = wb.create_sheet("Pivot")
    pivot["A1"] = "Region"
    pivot["B1"] = "North"
    pivot["A3"] = "Brand"
    pivot["B3"] = "加總 - Qty"
    pivot["A4"], pivot["B4"] = "BETA", 10
    pivot["A5"], pivot["B5"] = "ACME", 5
    pivot["A6"], pivot["B6"] = "總計", 15

    hidden = wb.create_sheet("工作表2")
    hidden["A1"], hidden["B1"] = "Key", "Value"
    hidden["A2"], hidden["B2"] = "secret", 1
    hidden.sheet_state = "hidden"
    wb.save(path)

    # openpyxl can't write pivot tables or x14 data validation (a list sourced from
    # another sheet, as Excel saves it): add a minimal pivot to "Pivot" and an
    # x14 validation to "Merged".
    pivot_index = wb.sheetnames.index("Pivot") + 1
    merged_name = f"xl/worksheets/sheet{wb.sheetnames.index('Merged') + 1}.xml"
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        rels_name = f"xl/worksheets/_rels/sheet{pivot_index}.xml.rels"
        for item in src.infolist():
            if item.filename == merged_name:
                xml = src.read(item.filename).decode("utf-8")
                dst.writestr(item, xml.replace("</worksheet>", X14_VALIDATION + "</worksheet>"))
            elif item.filename != rels_name:
                dst.writestr(item, src.read(item.filename))
        dst.writestr(
            rels_name,
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rIdPivot1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" '
            'Target="../pivotTables/pivotTable1.xml"/></Relationships>',
        )
        dst.writestr("xl/pivotTables/pivotTable1.xml", PIVOT_XML)
        dst.writestr("xl/pivotTables/_rels/pivotTable1.xml.rels", PIVOT_RELS)
        dst.writestr("xl/pivotCache/pivotCacheDefinition1.xml", CACHE_XML)
    tmp.replace(path)


def section(markdown: str, title: str) -> str:
    start = markdown.find(f"## {title}")
    if start < 0:
        return ""
    end = markdown.find("\n## ", start + 4)
    return markdown[start:end if end > 0 else len(markdown)]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cmd", help="command to test (default: the skill's convert.py)")
    args = parser.parse_args()
    if args.cmd:
        cmd = [p.strip('"') for p in shlex.split(args.cmd, posix=(sys.platform != "win32"))]
        if Path(cmd[0]).is_file():
            cmd[0] = str(Path(cmd[0]).resolve())
    else:
        cmd = DEFAULT_CMD

    failures = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  -- {detail[:400]}" if detail and not ok else ""))
        if not ok:
            failures.append(name)

    def run(*extra: str) -> str:
        proc = subprocess.run(cmd + [str(book), *extra], capture_output=True, timeout=300)
        if proc.returncode != 0:
            failures.append(f"exit code {proc.returncode} for {extra}")
            print(proc.stderr.decode("utf-8", "replace")[-1500:])
        return proc.stdout.decode("utf-8", "replace")

    with tempfile.TemporaryDirectory() as tmp:
        book = Path(tmp) / "planning report.xlsx"
        build_workbook(book)

        md = run()
        dns = section(md, "Report")
        # Item 4: real header row, week labels merged over dates, totals kept separately.
        check("overview names the header source", "第 13–14 列（依自動篩選範圍，合併上方標籤列）" in md, md[:800])
        check("header names from row 14", "| 列 | MPN (A) | Brand (B·隱) | Qty (C) |" in dns, dns[:1500])
        check("week label + date header", "W1(W38) 2026-09-21 (D)" in dns and "W2(W39) 2026-09-28 (E)" in dns)
        check("rows above the header kept apart", "### 表頭上方的列（第 1–13 列" in dns)
        check("no pandas placeholder names", "Unnamed" not in md and "| 0.1 |" not in md)
        # Item 5: Excel display formats, never scientific notation.
        check("accounting format keeps every digit", "| 12,345,678 [註1] |" in dns, dns[-2500:])
        check("negative in red-parentheses format", "(1,234)" in dns)
        check("accounting zero shows a dash", "| - |" in dns)
        check("General float without noise", "| 16.4 |" in dns)
        check("percent format", "206.92%" in dns)
        check("total row formatted", "1,234,568" in dns)
        check("no scientific notation", "e+0" not in md.lower())
        raw = section(run("--xlsx-raw-values"), "Report")
        check("raw values keep full precision", "| 12345678 [註1] |" in raw and "| 16.4 |" in raw and "2.06919466383172" in raw, raw[-1500:])
        # Item 6: hidden rows / columns / sheets, comments, validation, filter, merges.
        check("hidden row marked", "| 17 (隱) |" in dns)
        check("hidden sheet grouped and labelled", "# 隱藏的工作表" in md and "## 工作表2（隱藏工作表）" in md)
        check("comment marker and text", "[註1]" in dns and "請確認交期內的訂單是否足夠" in dns and "（Planner）" in dns)
        check("autofilter criteria", "A 欄：只顯示 PN-1001、PN-2002" in dns, dns[-1200:])
        check("data validation options", "清單，選項：OK、NG、TBD" in dns)
        check("freeze panes", "凍結窗格：前 14 列、前 1 欄" in dns)
        merged = section(md, "Merged")
        check("merged header cell fills the column name", "Date (B)" in merged and "B1:B2" in merged, merged)
        check("x14 data validation (list from another sheet)",
              "A3:A10、C3:C5、D3:D5、E3:E5 等 5 段：清單，選項來自 Report!$F$15:$F$17" in merged, merged)
        pivot = section(md, "Pivot")
        check("pivot gets its own block", "### 樞紐分析表「PivotFixture」（A3:B6）" in pivot, pivot)
        check("pivot header from its layout", "| 列 | Brand (A) | 加總 - Qty (B) |" in pivot)
        check("pivot layout and source", "列：Brand" in pivot and "值：加總 - Qty" in pivot and "'Report'!A14:E20" in pivot)
        check("pivot page filter", "目前的篩選：Region = North" in pivot)

        visible = run("--xlsx-visible-only")
        vdns = section(visible, "Report")
        check("visible-only drops hidden sheet", "工作表2" not in visible)
        check("visible-only drops hidden row and column", "HIDDEN-ROW" not in vdns and "Brand (B" not in vdns)

        override = section(run("--xlsx-header-row", "Report=1"), "Report")
        check("manual header row", "第 1 列（手動指定）" in override, override[:600])

        classic = run("--xlsx-classic")
        check("classic converter still available", "## Report" in classic and "Unnamed" in classic)

    print()
    if failures:
        print(f"{len(failures)} check(s) failed: {', '.join(failures)}")
        return 1
    print("All Excel checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
