#!/usr/bin/env python3
"""Test the compact Excel package for Claude (xlsx_compact.py, the default format).

Every value in the tab-separated data is compared with the workbook as read
independently by openpyxl, for a small workbook (single file) and a larger one
that is split into several files. Also checks the index (exact statistics,
formula logic), the keyword row filter, -o / --out-dir writing and cleanup.

    python tests/compact_test.py                              # the skill's convert.py
    python tests/compact_test.py --cmd dist/markitdown.exe    # the Windows build
"""

from __future__ import annotations

import argparse
import math
import random
import re
import shlex
import subprocess
import sys
import tempfile
import warnings
import zipfile
from datetime import datetime, time
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tests"))
sys.path.insert(0, str(ROOT / "skills" / "markitdown" / "scripts"))
DEFAULT_CMD = [sys.executable, str(ROOT / "skills" / "markitdown" / "scripts" / "convert.py")]
INDEX = "00_索引與邏輯.md"
CATEGORIES = ["Demand", "ETD", "Balance"]


def build_big(path: Path) -> None:
    from openpyxl import Workbook

    rng = random.Random(7)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["ID", "Name", "Amount", "Ratio", "Date", "Note", "Category", "", "", "Region"])
    for i in range(1, 3001):
        name = f"PART-{i:05d}"
        if i % 97 == 0:
            name += "\twith tab"
        if i % 89 == 0:
            name += "\nsecond line"
        if i % 83 == 0:
            name += " back\\slash | pipe"
        note = "空" if i % 5 else f"備註 {i}"
        ws.append([i, name, round(rng.uniform(-1e6, 1e6), 4), rng.random(),
                   datetime(2026, 1, 1 + i % 28), note if i % 7 else None, CATEGORIES[i % 3],
                   None, None, ["North", "South"][i // 50 % 2]])
    for r in range(10, 21):
        ws.row_dimensions[r].hidden = True
    # Formats that hide part of the value: a time of day behind a date-only format, a month-only format.
    ws["H1"], ws["I1"] = "Shipped", "Month"
    ws["H2"] = datetime(2026, 9, 24, 14, 22, 52, 999000)
    ws["H2"].number_format = "yyyy-mm-dd"
    ws["I2"] = datetime(2026, 10, 31)
    ws["I2"].number_format = "yyyy-mm"
    small = wb.create_sheet("Small")
    small.append(["k", "v"])
    small.append(["a", 1])
    wb.save(path)


def same_moment(field: str, value) -> bool:
    """Dates are written in full whatever the cell format shows: YYYY-MM-DD[ hh:mm:ss[.fff]]."""
    try:
        if isinstance(value, datetime):
            return datetime.fromisoformat(field) == value
        if isinstance(value, time):
            return time.fromisoformat(field) == value
    except ValueError:
        return False
    return False


def fifteen_digits(value: float) -> set:
    """Excel shows 15 significant digits; the stored decimal text is rounded half-up
    (the binary value may round the other way on an exact ...5)."""
    text = repr(float(value))
    exponent = Decimal(text).adjusted()
    half_up = Decimal(text).quantize(Decimal(1).scaleb(exponent - 14), rounding=ROUND_HALF_UP)
    return {float(f"{value:.15g}"), float(half_up)}


def same(field: str, value) -> bool:
    """A written field matches the workbook value (numbers to Excel's 15 significant digits)."""
    if isinstance(value, (datetime, time)):
        return same_moment(field, value)
    if isinstance(value, bool):
        return field.upper() == str(value).upper()
    if isinstance(value, (int, float)):
        try:
            return float(field) in fifteen_digits(value)
        except ValueError:
            return False
    return field == ("" if value is None else str(value))


def unescape(text: str) -> str:
    return re.sub(r"\\(.)", lambda m: {"t": "\t", "n": "\n", "r": "\r", "\\": "\\"}.get(m.group(1), m.group(0)), text)


def tsv_blocks(texts):
    """Yield (sheet, [lines]) for every data block in the given Markdown files."""
    for text in texts:
        for m in re.finditer(r"^## 資料：(.+?)(?:（第 \d+/\d+ 部分）)?\n.*?```tsv\n(.*?)\n```", text, re.S | re.M):
            yield m.group(1), m.group(2).split("\n")


def readable_copy(book: Path) -> Path:
    """openpyxl can't load the hand-made pivot table in xlsx_test's workbook: drop it."""
    copy = book.with_name(book.stem + ".openpyxl.xlsx")
    with zipfile.ZipFile(book) as src, zipfile.ZipFile(copy, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith(("xl/pivotTables/", "xl/pivotCache/")):
                continue
            if item.filename.startswith("xl/worksheets/_rels/") and b"pivotTable" in data:
                continue
            dst.writestr(item, data)
    return copy


def compare_with_openpyxl(book: Path, texts, check, label: str, only_rows=None) -> None:
    from openpyxl import load_workbook

    warnings.simplefilter("ignore")  # "Data Validation extension is not supported"
    wb = load_workbook(readable_copy(book), data_only=True)
    seen = {}
    mismatches = []
    for sheet, lines in tsv_blocks(texts):
        ws = wb[sheet]
        letters = lines[0].split("\t")[1:]
        for line in lines[1:]:
            fields = line.split("\t")
            label_text = fields[0]
            r = int(label_text.rstrip("h"))
            if label_text.endswith("h") != bool(ws.row_dimensions[r].hidden):
                mismatches.append(f"{sheet}!{r}: hidden marker")
            values = [unescape(f) for f in fields[1:]] + [""] * (len(letters) - len(fields) + 1)
            for letter, field in zip(letters, values):
                if not same(field, ws[f"{letter}{r}"].value):
                    mismatches.append(f"{sheet}!{letter}{r}: wrote {field!r}, openpyxl {ws[f'{letter}{r}'].value!r}")
            seen.setdefault(sheet, set()).add(r)
    for ws in wb.worksheets:
        rows = {c.row for row in ws.iter_rows() for c in row if c.value is not None}
        if only_rows is not None:
            rows = only_rows.get(ws.title, set())
        if rows != seen.get(ws.title, set()):
            missing = sorted(rows - seen.get(ws.title, set()))[:5]
            extra = sorted(seen.get(ws.title, set()) - rows)[:5]
            mismatches.append(f"{ws.title}: rows differ, missing {missing}, extra {extra}")
    check(f"{label}: every cell matches openpyxl", not mismatches, "; ".join(mismatches[:5]))


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cmd", help="command to test (default: the skill's convert.py)")
    args = parser.parse_args()
    if args.cmd:
        cmd = [p.strip('"') for p in shlex.split(args.cmd, posix=(sys.platform != "win32"))]
        if Path(cmd[0]).is_file():
            cmd[0] = str(Path(cmd[0]).resolve())
    else:
        cmd = DEFAULT_CMD

    from xlsx_test import build_workbook
    from xlsx_converter import general_text

    failures = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  -- {detail[:400]}" if detail and not ok else ""))
        if not ok:
            failures.append(name)

    def run(*argv: str, expect: int = 0) -> str:
        proc = subprocess.run(cmd + [str(a) for a in argv], capture_output=True, timeout=600)
        if proc.returncode != expect:
            failures.append(f"exit code {proc.returncode} for {argv}")
            print(proc.stderr.decode("utf-8", "replace")[-1500:])
        return proc.stdout.decode("utf-8", "replace")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        book = tmp / "planning report.xlsx"
        build_workbook(book)
        big = tmp / "big.xlsx"
        build_big(big)

        # 1. Small workbook: one file, index first, then the data.
        md = run(book)
        check("compact is the default", md.startswith("# planning report.xlsx（給 Claude 的精簡版）"), md[:200])
        check("single file holds index and data", "## 工作表" in md and "## 資料：Report" in md and "```tsv" in md)
        check("row numbers and hidden-row marker", "\n15\tPN-1001\tACME\t12345678\t-1234\t0\tOK" in md and "\n17h\tHIDDEN-ROW\t" in md)
        check("full precision, no display formatting", "\t2.06919466383172\t" in md and "12,345,678" not in md)
        check("dates in ISO form", "\t2026-09-21\t2026-09-28\t" in md)
        check("exact column total in the index", "| C | Qty | 數值 | 3 | 3 個數值，最小 1，最大 12345678，合計 12345695.4 |" in md)
        check("hidden column named in the index", "| B | Brand（隱藏欄） |" in md)
        check("formula logic per column", "| B | Times | 數值 | 3 |" in md and "3 格；`=A2*$D$2`（B2 的寫法）" in md)
        check("comments go with the data", "### 註解（Report）\n\n- C15（Planner）：請確認交期內的訂單是否足夠" in md)
        check("autofilter and validation in the index", "A 欄：只顯示 PN-1001、PN-2002" in md and "清單，選項：OK、NG、TBD" in md)
        check("data flow between sheets", "| Calc | Report（3 格） | - |" in md)
        compare_with_openpyxl(book, [md], check, "small workbook")

        # 2. Keyword filter keeps matching rows plus the header rows; statistics stay whole-sheet.
        filtered = run(book, "--xlsx-rows-containing", "pn-2002")
        report = next((lines for sheet, lines in tsv_blocks([filtered]) if sheet == "Report"), [])
        check("filter keeps header rows and matches", [line.split("\t")[0] for line in report[1:]] == ["13", "14", "16"], str(report))
        check("filter is announced", "已依關鍵字篩選" in filtered and "pn-2002" in filtered)
        check("statistics still cover the whole sheet", "合計 12345695.4" in filtered)
        check("sheets without matches are left out", "## 資料：Calc" not in filtered)
        run(book, "--xlsx-format", "readable", "--xlsx-rows-containing", "x", expect=2)

        # 3. Larger workbook split into several files in <name>_md/.
        out = tmp / "out"
        run(big, "--out-dir", out, "--xlsx-part-tokens", "4000")
        folder = out / "big_md"
        files = sorted(p.name for p in folder.glob("*.md"))
        check("split into a folder with the index first", files[:1] == [INDEX] and len(files) > 5, str(files))
        index = (folder / INDEX).read_text(encoding="utf-8")
        texts = [(folder / f).read_text(encoding="utf-8") for f in files if f != INDEX]
        check("index lists every data file", all(f"| {f} |" in index for f in files if f != INDEX))
        check("index stays small", len(index) < 6000, str(len(index)))
        check("parts named by sheet", any(f.startswith("01_Data_第") for f in files) and "02_Small.md" in files, str(files))
        check("every part repeats the column letters", all("```tsv\n列\tA\t" in t for t in texts))
        compare_with_openpyxl(big, texts, check, "split workbook")

        from openpyxl import load_workbook

        ws = load_workbook(big, data_only=True)["Data"]
        amounts = {c: [] for c in CATEGORIES}
        for row in ws.iter_rows(min_row=2, values_only=True):
            amounts[row[6]].append(row[2])
        expected = "、".join(f"{c} {general_text(math.fsum(amounts[c]), 15)}" for c in sorted(amounts, key=lambda c: CATEGORIES.index(c) if c != "Demand" else -1))
        check("totals split by the row category", "依分類合計" in index and all(
            f"{c} {general_text(math.fsum(v), 15)}" in index for c, v in amounts.items()), expected)
        check("row category found (G), plain grouping column (J) not", "列的分類：G 欄（Category）" in index
              and "列的分類：J" not in index, index[:3000])
        check("dates keep what the format hides", "\t2026-09-24 14:22:52.999\t2026-10-31" in "".join(texts))
        check("special characters escaped", "PART-00097\\twith tab" in "".join(texts) and "\\nsecond line" in "".join(texts))

        # 4. Re-running with bigger parts removes the old parts but nothing else.
        (folder / "my notes.txt").write_text("keep me", encoding="utf-8")
        run(big, "--out-dir", out, "--xlsx-part-tokens", "60000")
        files2 = sorted(p.name for p in folder.glob("*.md"))
        check("stale parts removed on re-run", len(files2) < len(files) and set(files2) <= set(files) | {"01_Data.md"}, str(files2))
        check("other files in the folder kept", (folder / "my notes.txt").exists())
        compare_with_openpyxl(big, [(folder / f).read_text(encoding="utf-8") for f in files2 if f != INDEX], check, "re-run")

        # 5. -o: a single file stays a file; a package becomes <name>_md/.
        single = tmp / "single.md"
        run(book, "-o", single)
        check("-o writes one file when it fits", single.exists() and "## 資料：Report" in single.read_text(encoding="utf-8"))
        target = tmp / "pkg.md"
        printed = run(big, "-o", target, "--xlsx-part-tokens", "4000")
        check("-o writes a package folder when split", (tmp / "pkg_md" / INDEX).exists() and "pkg_md" in printed, printed)

        # 6. To stdout, a split package prints the index only and says so.
        proc = subprocess.run(cmd + [str(big), "--xlsx-part-tokens", "4000"], capture_output=True, timeout=600)
        out_text = proc.stdout.decode("utf-8", "replace")
        check("stdout gets the index of a split package", proc.returncode == 0 and "## 資料檔" in out_text and "```tsv" not in out_text)
        check("stdout note explains how to get the data", "--out-dir" in proc.stderr.decode("utf-8", "replace"))

        # 7. The readable form is still available.
        readable = run(book, "--xlsx-format", "readable")
        check("readable format still available", "| 列 | MPN (A) | Brand (B·隱) | Qty (C) |" in readable)

    print()
    if failures:
        print(f"{len(failures)} check(s) failed: {', '.join(failures)}")
        return 1
    print("All compact-package checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
