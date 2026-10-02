"""
WPG Holdings 據點爬蟲
=====================

抓取 https://www.wpgholdings.com/branches/search/en/0_0_0/0/ 內所有公司據點資料，
輸出為 Excel 資料庫，欄位：

    Company | Region/City | Zip Code | Tel | Fax | Email | Address

用法：
    python scraper.py                         # 線上抓取，輸出 wpg_branches.xlsx
    python scraper.py -o out.xlsx             # 指定輸出檔名
    python scraper.py --browser               # 用 Playwright（頁面需 JavaScript 時）
    python scraper.py --html saved1.html ...  # 解析已存好的 HTML 檔（離線）
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from collections import deque
from dataclasses import dataclass, fields
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup, NavigableString, Tag
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

START_URL = "https://www.wpgholdings.com/branches/search/en/0_0_0/0/"
# 只跟隨同一個搜尋頁面底下的連結（分頁 / 篩選條件）
CRAWL_PREFIX = "/branches/search/en/"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

COLUMNS = ["Company", "Region/City", "Zip Code", "Tel", "Fax", "Email", "Address"]

# 標籤 -> 欄位。順序重要：較長 / 較特定的放前面。
LABELS = [
    (r"zip\s*code|postal\s*code|post\s*code|zip|郵遞區號|邮编|郵編", "zip_code"),
    (r"tel(?:ephone)?|phone|電話|电话", "tel"),
    (r"fax|傳真|传真", "fax"),
    (r"e-?mail|信箱|邮箱|郵箱", "email"),
    (r"address|addr\.?|add\.?|地址", "address"),
    (r"region\s*/\s*city|region|city|地區|地区|城市", "region_city"),
]
LABEL_RE = re.compile(
    r"^\s*(?P<label>" + "|".join(p for p, _ in LABELS) + r")\s*[:：.]?\s*(?P<value>.*)$",
    re.IGNORECASE,
)
TEL_LABEL_RE = re.compile(r"^\s*(tel(?:ephone)?|phone|電話|电话)\b", re.IGNORECASE)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
HEADING_TAGS = ["h1", "h2", "h3", "h4", "h5", "h6"]


@dataclass
class Branch:
    company: str = ""
    region_city: str = ""
    zip_code: str = ""
    tel: str = ""
    fax: str = ""
    email: str = ""
    address: str = ""

    def key(self) -> tuple:
        norm = lambda s: re.sub(r"\s+", " ", s).strip().lower()
        return (norm(self.company), norm(self.address), norm(self.tel))

    def row(self) -> list[str]:
        return [getattr(self, f.name) for f in fields(self)]


# --------------------------------------------------------------------------- #
# 解析
# --------------------------------------------------------------------------- #
def _field_for(label: str) -> str | None:
    for pattern, name in LABELS:
        if re.fullmatch(pattern, label.strip(), re.IGNORECASE):
            return name
    return None


def _clean(s: str) -> str:
    s = s.replace("\xa0", " ")
    s = re.sub(r"[ \t\r\f\v]+", " ", s)
    return s.strip(" :：\n")


def _lines(tag: Tag) -> list[str]:
    return [l for l in (_clean(x) for x in tag.get_text("\n").split("\n")) if l]


def _count_tel_labels(tag: Tag) -> int:
    return sum(1 for l in _lines(tag) if TEL_LABEL_RE.match(l))


def _find_record_blocks(soup: BeautifulSoup) -> list[Tag]:
    """找出每一筆據點的容器：包含剛好一個 'Tel' 標籤的最大祖先元素。"""
    blocks: list[Tag] = []
    seen: set[int] = set()
    for text in soup.find_all(string=TEL_LABEL_RE):
        node = text.parent
        if node is None or node.name in ("script", "style", "option", "title"):
            continue
        block = node
        while block.parent is not None and block.parent.name not in ("body", "html", "[document]"):
            if _count_tel_labels(block.parent) > 1:
                break
            block = block.parent
        if id(block) not in seen:
            seen.add(id(block))
            blocks.append(block)
    return blocks


def _nearest_heading_before(block: Tag) -> str:
    h = block.find_previous(HEADING_TAGS)
    return _clean(h.get_text(" ")) if h else ""


def parse_block(block: Tag) -> Branch:
    b = Branch()
    lines = _lines(block)
    pre_label: list[str] = []  # 第一個欄位標籤之前的文字（公司名、地區等）
    current: str | None = None
    seen_label = False

    for line in lines:
        m = LABEL_RE.match(line)
        name = _field_for(m.group("label")) if m else None
        if name:
            seen_label = True
            current = name
            value = _clean(m.group("value"))
            if value:
                old = getattr(b, name)
                setattr(b, name, f"{old} {value}".strip() if old else value)
            continue
        if not seen_label:
            pre_label.append(line)
        elif current:
            # 標籤和值分在不同元素 / 地址換行 → 接到目前欄位
            old = getattr(b, current)
            if current in ("tel", "fax", "zip_code", "email") and old:
                # 這些欄位只取一行，多出的文字不黏上去
                continue
            setattr(b, current, f"{old} {line}".strip() if old else line)

    # 公司名稱：優先使用區塊內的 heading / strong，否則取第一行
    title = block.find(HEADING_TAGS + ["strong", "b"])
    if title and not LABEL_RE.match(_clean(title.get_text(" "))):
        b.company = _clean(title.get_text(" "))
    elif pre_label:
        b.company = pre_label[0]

    # 地區/城市：區塊內標籤前的其他文字，或區塊前最近的標題
    if not b.region_city:
        others = [l for l in pre_label if l != b.company]
        if others:
            b.region_city = others[0]
        else:
            heading = _nearest_heading_before(block)
            if heading and heading != b.company:
                b.region_city = heading

    # Email：若有 mailto 連結或藏在其他地方，補抓
    if not EMAIL_RE.search(b.email or ""):
        a = block.find("a", href=re.compile(r"^mailto:", re.I))
        if a:
            b.email = a["href"].split(":", 1)[1].split("?")[0]
        else:
            m = EMAIL_RE.search(block.get_text(" "))
            b.email = m.group(0) if m else b.email
    else:
        b.email = EMAIL_RE.search(b.email).group(0)

    return b


def parse_html(html: str) -> list[Branch]:
    soup = BeautifulSoup(html, "html.parser")
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    return [parse_block(blk) for blk in _find_record_blocks(soup)]


def find_crawl_links(html: str, base_url: str) -> list[str]:
    soup = BeautifulSoup(html, "html.parser")
    host = urlparse(base_url).netloc
    out = []
    candidates = [a.get("href") for a in soup.find_all("a", href=True)]
    candidates += [o.get("value") for o in soup.find_all("option") if o.get("value")]
    for href in candidates:
        if not href or href.startswith(("javascript:", "mailto:", "tel:", "#")):
            continue
        url = urljoin(base_url, href).split("#")[0]
        p = urlparse(url)
        if p.netloc == host and p.path.startswith(CRAWL_PREFIX):
            out.append(normalize_url(url))
    return out


def normalize_url(url: str) -> str:
    p = urlparse(url)
    path = p.path if p.path.endswith("/") else p.path + "/"
    return p._replace(path=path, fragment="").geturl()


# --------------------------------------------------------------------------- #
# 抓取
# --------------------------------------------------------------------------- #
class RequestsFetcher:
    def __init__(self, delay: float):
        self.s = requests.Session()
        self.s.headers.update(HEADERS)
        self.delay = delay

    def get(self, url: str) -> str:
        for attempt in range(4):
            try:
                r = self.s.get(url, timeout=30)
                r.raise_for_status()
                r.encoding = r.apparent_encoding or r.encoding
                time.sleep(self.delay)
                return r.text
            except requests.RequestException as e:
                if attempt == 3:
                    raise
                wait = 2 ** (attempt + 1)
                print(f"  ! {e} — {wait}s 後重試", file=sys.stderr)
                time.sleep(wait)
        return ""

    def close(self):
        self.s.close()


class BrowserFetcher:
    """頁面資料由 JavaScript 載入時使用（pip install playwright && playwright install chromium）。"""

    def __init__(self, delay: float):
        from playwright.sync_api import sync_playwright

        self._pw = sync_playwright().start()
        self._browser = self._pw.chromium.launch()
        self._page = self._browser.new_page(user_agent=HEADERS["User-Agent"])
        self.delay = delay

    def get(self, url: str) -> str:
        self._page.goto(url, wait_until="networkidle", timeout=60000)
        time.sleep(self.delay)
        return self._page.content()

    def close(self):
        self._browser.close()
        self._pw.stop()


def crawl(start_url: str, fetcher, max_pages: int) -> list[Branch]:
    queue = deque([normalize_url(start_url)])
    visited: set[str] = set()
    results: dict[tuple, Branch] = {}

    while queue and len(visited) < max_pages:
        url = queue.popleft()
        if url in visited:
            continue
        visited.add(url)
        try:
            html = fetcher.get(url)
        except Exception as e:  # noqa: BLE001
            print(f"[skip] {url}: {e}", file=sys.stderr)
            continue

        branches = parse_html(html)
        new = 0
        for b in branches:
            if b.key() not in results:
                results[b.key()] = b
                new += 1
        print(f"[{len(visited):>3}] {url}  找到 {len(branches)} 筆，新增 {new} 筆（累計 {len(results)}）")

        for link in find_crawl_links(html, url):
            if link not in visited:
                queue.append(link)

    return list(results.values())


# --------------------------------------------------------------------------- #
# Excel 輸出
# --------------------------------------------------------------------------- #
def write_excel(branches: list[Branch], path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Branches"
    ws.append(COLUMNS)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for b in sorted(branches, key=lambda x: (x.region_city.lower(), x.company.lower())):
        ws.append(b.row())

    widths = {"Company": 40, "Region/City": 20, "Zip Code": 12, "Tel": 22,
              "Fax": 22, "Email": 34, "Address": 70}
    for i, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths[col]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.number_format = "@"  # 以文字儲存，保留郵遞區號 / 電話前導 0

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="WPG Holdings 據點爬蟲 → Excel")
    ap.add_argument("-o", "--output", default="wpg_branches.xlsx", help="輸出 Excel 檔名")
    ap.add_argument("--url", default=START_URL, help="起始網址")
    ap.add_argument("--max-pages", type=int, default=300, help="最多抓取頁數")
    ap.add_argument("--delay", type=float, default=1.0, help="每頁間隔秒數")
    ap.add_argument("--browser", action="store_true", help="使用 Playwright 渲染 JavaScript")
    ap.add_argument("--html", nargs="+", metavar="FILE", help="改為解析本機 HTML 檔")
    args = ap.parse_args(argv)

    if args.html:
        seen: dict[tuple, Branch] = {}
        for path in args.html:
            with open(path, encoding="utf-8", errors="replace") as f:
                for b in parse_html(f.read()):
                    seen.setdefault(b.key(), b)
        branches = list(seen.values())
    else:
        fetcher = BrowserFetcher(args.delay) if args.browser else RequestsFetcher(args.delay)
        try:
            branches = crawl(args.url, fetcher, args.max_pages)
        finally:
            fetcher.close()

    if not branches:
        print("沒有抓到任何資料。若網頁需 JavaScript，請改用 --browser。", file=sys.stderr)
        return 1

    write_excel(branches, args.output)
    print(f"完成：共 {len(branches)} 筆，已輸出至 {args.output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
