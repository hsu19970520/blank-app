#!/usr/bin/env python3
"""Tests for "recalculate with Excel before converting" (excel_recalc.py).

Real Excel isn't available in CI, so a fake Excel object stands in for COM and
records every call. The fake "recalculates" by changing a cached value in the
copy, which proves the Markdown comes from the recalculated copy while the
original file stays untouched. A last check runs the CLI with --xlsx-recalc
on a machine without Excel and expects a clear refusal.

    python tests/recalc_test.py
    python tests/recalc_test.py --cmd dist/markitdown.exe
"""

from __future__ import annotations

import argparse
import hashlib
import shlex
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "skills" / "markitdown" / "scripts"
sys.path.insert(0, str(SCRIPTS))

import excel_recalc  # noqa: E402


class FakeExcel:
    """Just enough of Excel.Application for excel_recalc.recalculate()."""

    def __init__(self, log, fail_open=False, name_error=False):
        self.__dict__.update(log=log, fail_open=fail_open, name_error=name_error, Version="16.0", settings={})
        self.__dict__["Workbooks"] = self

    def __setattr__(self, name, value):
        self.settings[name] = value

    def Open(self, path, **kwargs):
        self.log.append(("Open", Path(path).name, kwargs.get("UpdateLinks")))
        if self.fail_open:
            raise RuntimeError("檔案已損毀")
        return FakeWorkbook(self.log, Path(path), self.name_error)

    def CalculateFull(self):
        self.log.append(("CalculateFull",))

    def Quit(self):
        self.log.append(("Quit",))


class FakeWorkbook:
    def __init__(self, log, path, name_error):
        self.log, self.path, self.name_error = log, path, name_error
        self.Worksheets = [FakeSheet(2), FakeSheet(0), FakeSheet(1)]

    def PivotCaches(self):
        return FakeCaches(self.log)

    def Save(self):
        self.log.append(("Save",))
        import openpyxl  # "recalculate": change a cached value in the copy

        wb = openpyxl.load_workbook(self.path)
        wb.active["B2"] = "#NAME?" if self.name_error else 999
        wb.save(self.path)
        if self.name_error:  # make it an error cell the way Excel stores it
            _rewrite_as_error(self.path)

    def Close(self, SaveChanges=None):
        self.log.append(("Close", SaveChanges))


class FakeSheet:
    def __init__(self, pivots):
        self.pivots = pivots

    def PivotTables(self):
        return type("PivotTables", (), {"Count": self.pivots})()


class FakeCaches:
    Count = 2

    def __init__(self, log):
        self.log = log

    def Item(self, i):
        log = self.log

        class Cache:
            def Refresh(self):
                log.append(("Refresh", i))
                if i == 2:
                    raise RuntimeError("找不到資料來源")

        return Cache()


def _rewrite_as_error(path: Path) -> None:
    import re
    import zipfile

    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                data = re.sub(rb'<c r="B2"([^>]*)t="(?:s|inlineStr)"([^>]*)>.*?</c>',
                              rb'<c r="B2"\1t="e"\2><v>#NAME?</v></c>', data, flags=re.S)
            dst.writestr(item, data)
    tmp.replace(path)


def make_workbook(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Item", "Qty"])
    ws.append(["A", 1])
    wb.save(path)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--cmd", help="CLI to test the no-Excel refusal with (default: the skill's convert.py)")
    args = parser.parse_args()
    failures = []

    def check(name: str, ok: bool, detail: str = "") -> None:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  -- {detail}" if detail and not ok else ""))
        if not ok:
            failures.append(name)

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        book = tmp / "報表 W39.xlsx"
        make_workbook(book)
        before = digest(book)

        # 1. Successful run: order of operations, settings, cleanup, result note.
        log = []
        excels = []

        def dispatch(prog_id):
            check("dispatches Excel.Application", prog_id == "Excel.Application", prog_id)
            excels.append(FakeExcel(log))
            return excels[-1]

        messages = []
        result = excel_recalc.recalculate(book, dispatch=dispatch, progress=messages.append)
        steps = [entry[0] for entry in log]
        check("calculate, refresh pivots, calculate again, save, close, quit",
              steps == ["Open", "CalculateFull", "Refresh", "Refresh", "CalculateFull", "Save", "Close", "Quit"], str(steps))
        check("links are not updated", log[0][2] == 0, str(log[0]))
        check("Excel runs hidden, silent, with macros disabled",
              excels[0].settings.get("Visible") is False and excels[0].settings.get("DisplayAlerts") is False
              and excels[0].settings.get("AutomationSecurity") == 3, str(excels[0].settings))
        check("workbook closed without saving again", ("Close", False) in log)
        check("works on a copy; original untouched", result.path != book and digest(book) == before)
        check("copy keeps the file name", result.path.name == book.name)
        check("pivot tables counted", result.pivot_tables == 3, str(result.pivot_tables))
        check("a failed pivot refresh is reported", len(result.pivot_failures) == 1 and "找不到資料來源" in result.pivot_failures[0])
        note = result.note()
        check("note names Excel and the pivot refresh",
              "Microsoft Excel 16.0" in note and "重新整理 3 個樞紐分析表" in note and "1 個樞紐分析表資料無法重新整理" in note, note)
        check("progress messages in Chinese", any("重新計算" in m for m in messages), str(messages))
        excel_recalc_workdir = result.workdir
        check("temporary folder exists until the caller removes it", excel_recalc_workdir.exists())

        # 2. Conversion reads the recalculated copy (value 1 -> 999) and says so.
        import convert

        real = excel_recalc.recalculate
        excel_recalc.recalculate = lambda path, progress=None, **kw: real(path, dispatch=lambda _: FakeExcel([]), progress=progress)
        try:
            converter = convert.build_converter()
            md = convert.convert_source(converter, convert.Source(str(book)), xlsx_recalc=True)
        finally:
            excel_recalc.recalculate = real
        check("Markdown uses the recalculated value", "| 2 | A | 999 |" in md, md[-600:])
        check("Markdown says it was recalculated", "已用 Microsoft Excel 16.0 在轉換前重新計算" in md, md[:600])
        check("original still untouched after conversion", digest(book) == before)
        leftovers = [p for p in Path(tempfile.gettempdir()).glob("markitdown-recalc-*") if p != excel_recalc_workdir]
        check("conversion removes its temporary copy", not leftovers, str(leftovers))

        # 3. New #NAME? errors (Excel without XLOOKUP, say) are flagged.
        result = excel_recalc.recalculate(book, dispatch=lambda _: FakeExcel([], name_error=True))
        check("new #NAME? errors detected", result.new_name_errors == 1, str(result.new_name_errors))
        check("note warns the numbers are unreliable", "#NAME?" in result.note() and "不可信" in result.note())
        import shutil

        shutil.rmtree(result.workdir, ignore_errors=True)  # the caller's job after converting

        # 4. A failure still quits Excel and cleans up, with a readable message.
        log = []
        try:
            excel_recalc.recalculate(book, dispatch=lambda _: FakeExcel(log, fail_open=True))
            check("open failure raises RecalcError", False)
        except excel_recalc.RecalcError as exc:
            check("open failure raises RecalcError", "Excel 重新計算失敗：檔案已損毀" in str(exc), str(exc))
        check("Excel quit after the failure", log[-1] == ("Quit",), str(log))
        leftovers = [p for p in Path(tempfile.gettempdir()).glob("markitdown-recalc-*") if p != excel_recalc_workdir]
        check("failed run leaves no temporary copy", not leftovers, str(leftovers))
        shutil.rmtree(excel_recalc_workdir, ignore_errors=True)

        # 5. Without Excel the CLI refuses clearly instead of converting stale numbers silently.
        available, reason = excel_recalc.excel_status()
        if available:
            print("SKIP  CLI refusal (Excel is installed on this machine)")
        else:
            if args.cmd:
                cmd = [part.strip('"') for part in shlex.split(args.cmd, posix=(sys.platform != "win32"))]
                if Path(cmd[0]).is_file():
                    cmd[0] = str(Path(cmd[0]).resolve())
            else:
                cmd = [sys.executable, str(SCRIPTS / "convert.py")]
            proc = subprocess.run(cmd + [str(book), "--xlsx-recalc"], capture_output=True, timeout=300)
            err = proc.stderr.decode("utf-8", "replace")
            check("CLI without Excel exits 2 with the reason", proc.returncode == 2 and "--xlsx-recalc：" in err and reason in err,
                  f"exit {proc.returncode}: {err[-400:]}")

    print()
    if failures:
        print(f"{len(failures)} check(s) failed: {', '.join(failures)}")
        return 1
    print("All recalculation checks passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
